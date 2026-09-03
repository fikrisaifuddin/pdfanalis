import re
from django.contrib import messages
from django.shortcuts import render

from openpyxl import load_workbook

from .forms import WaBlastUploadForm, DEFAULT_TEMPLATE

NAMA_ALIASES = {"nama", "name", "nama lengkap"}
HP_ALIASES = {"no hp", "no. hp", "nomor hp", "no telp", "no. telp", "whatsapp", "wa", "hp", "nomor"}


def _cari_index_kolom(header_row):
    idx_nama, idx_hp = None, None
    for i, cell in enumerate(header_row):
        label = str(cell.value or "").strip().lower()
        if label in NAMA_ALIASES:
            idx_nama = i
        elif label in HP_ALIASES:
            idx_hp = i
    return idx_nama, idx_hp


def _format_nomor_wa(raw):
    if raw is None:
        return None

    if isinstance(raw, (int, float)):
        raw = str(int(raw))

    digits = re.sub(r"\D", "", str(raw))
    if not digits:
        return None
    if digits.startswith("0"):
        digits = "62" + digits[1:]
    elif digits.startswith("8"):
        digits = "62" + digits
    if not digits.startswith("62") or len(digits) < 9:
        return None
    return digits


def index(request):
    context = {
        "form": WaBlastUploadForm(),
        "kontak_list": None,
    }

    if request.method == "POST":
        form = WaBlastUploadForm(request.POST, request.FILES)
        context["form"] = form

        if form.is_valid():
            template_text = form.cleaned_data["template"] or DEFAULT_TEMPLATE
            excel_file = form.cleaned_data["excel_file"]

            try:
                wb = load_workbook(excel_file, data_only=True)
                sheet = wb.active
                rows = list(sheet.iter_rows())
            except Exception:
                messages.error(request, "File Excel tidak bisa dibaca. Pastikan formatnya .xlsx yang valid.")
                return render(request, "wablast/wablast.html", context)

            if not rows:
                messages.error(request, "File Excel kosong.")
                return render(request, "wablast/wablast.html", context)

            idx_nama, idx_hp = _cari_index_kolom(rows[0])
            if idx_nama is None or idx_hp is None:
                messages.error(
                    request,
                    "Kolom 'Nama' dan/atau 'No HP' tidak ditemukan di baris pertama file Excel. "
                    "Pastikan ada header kolom dengan nama tersebut.",
                )
                return render(request, "wablast/wablast.html", context)

            kontak_list = []
            for row_num, row in enumerate(rows[1:], start=2):
                nama_cell = row[idx_nama].value if idx_nama < len(row) else None
                hp_cell = row[idx_hp].value if idx_hp < len(row) else None

                nama = str(nama_cell).strip() if nama_cell else ""
                nomor_formatted = _format_nomor_wa(hp_cell)

                if not nama and not hp_cell:
                    continue

                # {sapaan} SENGAJA belum diganti di sini, biarkan tetap "{sapaan}"
                # supaya bisa diisi belakangan (di browser) sesuai dropdown per baris.
                pesan_raw = template_text.replace("{nama}", nama or "-")

                kontak_list.append(
                    {
                        "row_num": row_num,
                        "nama": nama or "(tanpa nama)",
                        "no_hp_asli": hp_cell,
                        "no_hp_formatted": nomor_formatted,
                        "pesan_raw": pesan_raw,
                        "valid": nomor_formatted is not None,
                    }
                )

            if not kontak_list:
                messages.warning(request, "Tidak ada data kontak yang terbaca dari file Excel.")

            context["kontak_list"] = kontak_list
            context["template_text"] = template_text

    return render(request, "wablast/wablast.html", context)